# -*- coding: utf-8 -*-
"""Fill the blanks of the template docx with answers (template structure unchanged)."""
import shutil
import docx
from docx.shared import Pt

SRC = r"C:\Users\Vincent\Downloads\TUGAS KELAS A.docx"
OUT = r"C:\Users\Vincent\Downloads\TUGAS KELAS A - JAWABAN.docx"

shutil.copy(SRC, OUT)
doc = docx.Document(OUT)

ANSWERS = [
    # (nomor, jawaban) - sesuai tabel 1..6
    (1,
     "Yang terjadi adalah method Aktivitas() yang dijalankan adalah versi override dari kelas "
     "masing-masing, yaitu versi override kelas Dokter (menampilkan \"Dokter ... sedang memeriksa "
     "pasien di ruang praktik\") dan versi override kelas Pasien (menampilkan \"... sedang "
     "beristirahat di kamar rumah sakit\"), bukan versi dari kelas Orang.\n"
     "Mengapa demikian: karena method Aktivitas() pada kelas Orang dideklarasikan dengan keyword "
     "virtual, lalu di-override di setiap subclass, sehingga C# melakukan dynamic dispatch "
     "(late binding), yaitu memilih method berdasarkan tipe objek sebenarnya saat runtime, bukan "
     "tipe variabelnya. Hal ini disebut polymorphism (banyak rupa) dalam PBO."),
    (2,
     "Method Diagnosa() berjalan normal dan menampilkan informasi bahwa dokter sedang "
     "mendiagnosa pasien.\n"
     "Mengapa demikian: karena method Diagnosa() didefinisikan langsung di dalam kelas Dokter, "
     "dan objek yang digunakan bertipe Dokter. Karena tipe variabel sesuai dengan kelas tempat "
     "method didefinisikan, method tersebut dapat dipanggil langsung. Jika objek disimpan dalam "
     "variabel bertipe Orang atau TenagaMedis, maka Diagnosa() tidak bisa dipanggil tanpa "
     "downcast, misalnya: ((Dokter)variabel).Diagnosa(...)."),
    (3,
     "Ya, bisa. Caranya sebagai berikut:\n"
     "1. Memanggil method InfoOrang() yang diwarisi dari kelas Orang untuk menampilkan nama dan "
     "umur, lalu memanggil CekSpesialis() yang diwarisi dari kelas TenagaMedis untuk menampilkan "
     "spesialisasi.\n"
     "2. Mengakses properti secara langsung: dokter.Nama, dokter.Umur, dan dokter.Spesialisasi.\n"
     "3. Jika objek disimpan dalam variabel bertipe Orang, gunakan downcast, misalnya "
     "((TenagaMedis)orang).CekSpesialis().\n"
     "4. Alternatif lain: membuat method baru, misalnya InfoLengkap(), di kelas Dokter yang "
     "menggabungkan seluruh informasi tersebut."),
    (4,
     "Method Menangis() berjalan normal dan menampilkan aktivitas menangis dari pasien anak.\n"
     "Mengapa demikian: karena method Menangis() didefinisikan di kelas PasienAnak, dan objek "
     "yang digunakan bertipe PasienAnak, sehingga objek langsung memiliki akses ke method "
     "tersebut. Jika objek disimpan dalam variabel bertipe Pasien atau Orang, method Menangis() "
     "tidak bisa dipanggil tanpa downcast ke PasienAnak."),
    (5,
     "Yang terjadi adalah yang dijalankan adalah Aktivitas() versi Perawat (\"Perawat ... sedang "
     "merawat dan memantau pasien\"), BUKAN versi kelas Orang.\n"
     "Mengapa demikian: variabel bertipe Orang diperbolehkan menyimpan objek subclass apa pun "
     "(upcasting). Karena Aktivitas() bersifat virtual dan telah di-override oleh kelas Perawat, "
     "C# melakukan dynamic dispatch: pemilihan method ditentukan oleh tipe objek sebenarnya saat "
     "runtime (Perawat), bukan tipe variabel saat compile time (Orang).\n"
     "Catatan: method khusus Perawat seperti CekPasien() tidak bisa dipanggil langsung dari "
     "variabel bertipe Orang; harus menggunakan downcast: ((Perawat)orang).CekPasien(...)."),
    (6,
     "1. Inheritance (pewarisan) memungkinkan subclass menggunakan kembali properti dan method "
     "dari superclass, sehingga program lebih ringkas, terstruktur, dan mudah dikembangkan.\n"
     "2. Polymorphism (virtual + override) membuat method yang sama dapat berperilaku berbeda "
     "sesuai objek sebenarnya, meskipun dipanggil melalui variabel bertipe induk. Contohnya "
     "List<Orang> dapat menampung objek Dokter, Perawat, dan Pasien sekaligus.\n"
     "3. Method khusus subclass hanya dapat diakses bila tipe variabel sesuai dengan kelas "
     "tempat method itu didefinisikan, atau dengan menggunakan downcast.\n"
     "4. Pewarisan bertingkat (multilevel inheritance) memperjelas hierarki sistem: "
     "Orang -> TenagaMedis -> Dokter/Perawat, dan Orang -> Pasien -> PasienAnak/PasienDewasa."),
]


def fill_cell(cell, text):
    """Put text into the single empty run of the cell's paragraph, matching template font."""
    p = cell.paragraphs[0]
    if not p.runs:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
    else:
        r = p.runs[0]
        r.text = text
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)


def fill_identity(doc, label, value):
    """Append value to the identity paragraph that starts with 'label :'."""
    for p in doc.paragraphs:
        if p.text.strip().startswith(label + " :"):
            p.runs[0].text = f"{label} : {value}"
            print(f"Filled identity: {label} : {value}")
            return True
    raise KeyError(f"identity paragraph for {label!r} not found")


fill_identity(doc, "Nama", "Vincent Alexander")
fill_identity(doc, "NIM", "252410103018")
fill_identity(doc, "Link Repository GitHub", "https://github.com/GylaxDev/A_252410103018_TASKPBO")
# Kelas sudah terisi 'A' pada template -> tidak diubah

for i, (n, ans) in enumerate(ANSWERS):
    table = doc.tables[i]
    # verify this table matches the expected question number
    first_col = table.rows[0].cells[0].text.strip()
    assert first_col == str(n), f"table {i} mismatch: {first_col}"
    answer_cell = table.rows[1].cells[2]
    fill_cell(answer_cell, ans)
    print(f"Soal {n}: jawaban terisi ({len(ans)} chars)")

doc.save(OUT)
print("Saved:", OUT)
